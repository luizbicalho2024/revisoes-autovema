"""Identificação conservadora de revisões no catálogo e nas OS ROF080.

Nunca converte quilometragem em ordinal sem uma regra homologada.
"""
from __future__ import annotations

import re
import unicodedata
from typing import Any

# Códigos não pessoais presentes no arquivo anomalias_revisao.xlsx da operação.
DEFAULT_CATALOG = [
    ("0010C21", "1ª REVISÃO FLEX - KM OU TEMPO"),
    ("0010C22", "1ª REVISÃO DIESEL - KM OU TEMPO"),
    ("0010F71", "1ª REVISÃO TITANO DIESEL - KM OU TEMPO"),
    ("0010F36", "2ª REVISÃO FLEX - KM OU TEMPO"),
    ("0010F37", "3ª REVISÃO FLEX - KM OU TEMPO"),
    ("0010F38", "4ª REVISÃO FLEX / 2ª REVISÃO DIESEL - KM OU TEMPO"),
    ("0010F39", "5ª REVISÃO FLEX - KM OU TEMPO"),
    ("0010F41", "6ª REVISÃO FLEX / 3ª REVISÃO DIESEL - KM OU TEMPO"),
    ("0010F42", "7ª REVISÃO FLEX - KM OU TEMPO"),
    ("0010F43", "8ª REVISÃO FLEX / 4ª REVISÃO DIESEL - KM OU TEMPO"),
    ("0010F44", "9ª REVISÃO FLEX - KM OU TEMPO"),
    ("0010F46", "10ª REVISÃO FLEX / 5ª REVISÃO DIESEL - KM OU TEMPO"),
    ("0010F47", "11ª REVISÃO FLEX - KM OU TEMPO"),
    ("0010F48", "12ª REVISÃO FLEX / 6ª REVISÃO DIESEL - KM OU TEMPO"),
    ("0010F49", "13ª REVISÃO FLEX - KM OU TEMPO"),
    ("0010F51", "14ª REVISÃO FLEX / 7ª REVISÃO DIESEL - KM OU TEMPO"),
    ("0010F52", "15ª REVISÃO FLEX - KM OU TEMPO"),
    ("0010F53", "16ª REVISÃO FLEX / 8ª REVISÃO DIESEL - KM OU TEMPO"),
    ("0010F54", "17ª REVISÃO FLEX - KM OU TEMPO"),
    ("0010F55", "18ª REVISÃO FLEX / 9ª REVISÃO DIESEL - KM OU TEMPO"),
    ("0010F56", "10ª REVISÃO DIESEL - KM OU TEMPO"),
    ("0010F57", "11ª REVISÃO DIESEL - KM OU TEMPO"),
    ("0010F58", "12ª REVISÃO DIESEL - KM OU TEMPO"),
]


def normalize(value: Any) -> str:
    plain = unicodedata.normalize("NFKD", str(value or ""))
    plain = "".join(c for c in plain if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", plain.upper()).strip()


def fuel_for_model(model: Any) -> str | None:
    text = normalize(model)
    if "DIESEL" in text or "TITANO" in text:
        return "diesel"
    if "FLEX" in text or "GASOLINA" in text or "ETANOL" in text:
        return "flex"
    return None


def possible_revision_numbers(text: Any, model: Any = None) -> list[int]:
    s = normalize(text)
    # Suporta "1A. REVISAO", "4A.REV.FLEX", "2° revisão" e variantes.
    found = [int(m.group(1)) for m in re.finditer(
        r"\b(\d{1,2})\s*(?:A|O|ª|º|°)?\s*\.?\s*REV(?:ISAO|\.|\b)", s
    )]
    words = {
        "PRIMEIRA": 1, "SEGUNDA": 2, "TERCEIRA": 3,
        "QUARTA": 4, "QUINTA": 5, "SEXTA": 6,
        "SETIMA": 7, "OITAVA": 8, "NONA": 9, "DECIMA": 10,
    }
    for word, n in words.items():
        if re.search(rf"\b{word}\s+REVISAO\b", s):
            found.append(n)
    if not found:
        return []
    # Em códigos mistos (ex.: 4ª FLEX / 2ª DIESEL) resolve pelo modelo.
    if len(set(found)) > 1 and "FLEX" in s and "DIESEL" in s:
        fuel = fuel_for_model(model)
        if fuel == "flex":
            matched = re.search(r"(\d{1,2})\s*(?:A|O|ª|º|°)?\s*\.?\s*REV[^/\n]{0,40}FLEX", s)
            if matched:
                return [int(matched.group(1))]
            return [found[0]]
        if fuel == "diesel":
            matched = re.search(r"(\d{1,2})\s*(?:A|O|ª|º|°)?\s*\.?\s*REV[^/\n]{0,40}DIESEL", s)
            if matched:
                return [int(matched.group(1))]
            return [found[-1]]
    return sorted(set(n for n in found if 1 <= n <= 30))


def read_catalog(file_bytes: bytes) -> list[dict[str, str]]:
    from io import BytesIO
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        ws = workbook["Anomalias"] if "Anomalias" in workbook.sheetnames else workbook.active
        iterator = ws.iter_rows(values_only=True)
        headers = [normalize(h) for h in next(iterator, [])]
        mandatory = ("CODIGO LOCAL", "DESCRICAO")
        if any(h not in headers for h in mandatory):
            raise ValueError("Catálogo inválido: esperadas as colunas CÓDIGO LOCAL e DESCRIÇÃO.")
        ci, di = headers.index("CODIGO LOCAL"), headers.index("DESCRICAO")
        result = []
        for row in iterator:
            code = str(row[ci] or "").strip().upper() if ci < len(row) else ""
            description = str(row[di] or "").strip() if di < len(row) else ""
            if re.fullmatch(r"[A-Z0-9]{3,24}", code) and description:
                result.append({"code": code, "description": description})
        if not result:
            raise ValueError("Nenhum código de revisão válido encontrado.")
        return result
    finally:
        workbook.close()


def catalog_from_database(db: Any) -> dict[str, str]:
    data = {code: desc for code, desc in DEFAULT_CATALOG}
    for row in db.revision_catalog.find({}, {"_id": 1, "description": 1}):
        data[str(row["_id"]).upper()] = str(row.get("description") or "")
    return data


def parse_evidence(service: Any, description: Any, observation: Any, model: Any,
                   catalog: dict[str, str]) -> dict[str, Any] | None:
    service_text = normalize(service)
    desc_text = normalize(description)
    if not service_text and not desc_text and not observation:
        return None
    source = "servico"
    original = service_text
    if service_text in catalog:
        original = normalize(catalog[service_text])
        source = "codigo"
    numbers = possible_revision_numbers(original, model)
    if not numbers and desc_text:
        numbers = possible_revision_numbers(desc_text, model)
    if numbers:
        fuel = fuel_for_model(model)
        text_flex = "FLEX" in original
        text_diesel = "DIESEL" in original
        conflict = bool((fuel == "diesel" and text_flex and not text_diesel)
                        or (fuel == "flex" and text_diesel and not text_flex)
                        or (fuel is None and text_flex and text_diesel and len(numbers) > 1))
        return {
            "revision_number": numbers[0] if len(numbers) == 1 and not conflict else None,
            "confidence": "servico" if len(numbers) == 1 and not conflict else "ambiguo",
            "evidence_source": source,
            "evidence_text": (service_text or desc_text)[:140],
        }
    # Revisão por quilometragem sem ordinal não deve ser transformada
    # em 1ª/2ª por pressuposto sobre periodicidade de cada modelo.
    service_revision = bool(re.search(r"\bREVISAO\b", service_text + " " + desc_text))
    if service_revision and not re.search(r"BOMBA|CABECOTE|INJECAO|CORREIA|SISTEMA", service_text):
        km = re.search(r"\b(\d{1,3})[. ]?(\d{3})\s*KM\b", service_text + " " + desc_text)
        return {
            "revision_number": None, "confidence": "sem_numero",
            "evidence_source": "servico", "evidence_text": service_text[:140],
            "reference_km": int(km.group(1) + km.group(2)) if km else None,
        }
    # Observação é evidência declaratória, não comprovação de execução.
    note_numbers = possible_revision_numbers(observation, model)
    if note_numbers:
        return {
            "revision_number": note_numbers[0] if len(note_numbers) == 1 else None,
            "confidence": "observacao" if len(note_numbers) == 1 else "ambiguo",
            "evidence_source": "observacao", "evidence_text": "Menção a revisão na observação da OS",
        }
    if re.search(r"\bREVISAO\b", normalize(observation)):
        return {
            "revision_number": None, "confidence": "observacao_sem_numero",
            "evidence_source": "observacao", "evidence_text": "Menção a revisão sem etapa identificada",
        }
    return None
